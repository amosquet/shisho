"""
utils/excel_manager.py - Core Excel (.xlsx) spreadsheet engine for Shisho.

Provides secure path resolution in data/excel/, workbook CRUD, formula preservation,
sheet management, range reading/writing, aligned Markdown table formatting with
Discord message truncation protection, file retention purging, and AI workbook summaries.
"""

from __future__ import annotations

import io
import logging
import os
import re
import time
import xml.etree.ElementTree
import zipfile
from collections.abc import Generator, Sequence
from contextlib import contextmanager
from pathlib import Path
from typing import Any, Union

import openpyxl
from openpyxl import Workbook
from openpyxl.cell import Cell
from openpyxl.utils.cell import (
    COORD_RE,
    coordinate_to_tuple,
    get_column_letter,
    range_boundaries,
)
from openpyxl.utils.exceptions import InvalidFileException
from openpyxl.worksheet.worksheet import Worksheet

logger = logging.getLogger(__name__)

# Repository root and default storage directory
REPO_ROOT = Path(__file__).resolve().parent.parent
DEFAULT_EXCEL_DIR = (REPO_ROOT / "data" / "excel").resolve()

# Excel worksheet title constraints
MAX_SHEET_TITLE_LEN = 31
FORBIDDEN_SHEET_CHARS = {":", "\\", "/", "?", "*", "[", "]"}
FORBIDDEN_TITLE_REGEX = re.compile(r"[\\*?:/\[\]]")

# Type aliases
WorkbookInput = Union[Workbook, str, Path]

# Exceptions raised when an Excel archive or XML component is damaged or corrupted
CORRUPT_EXCEL_EXCEPTIONS = (
    InvalidFileException,
    zipfile.BadZipFile,
    KeyError,
    OSError,
    xml.etree.ElementTree.ParseError,
)


# =========================================================================
# Path Resolution & Sandboxing
# =========================================================================

def get_excel_dir() -> Path:
    """
    Returns the absolute Path to the data/excel storage directory,
    creating it if it does not already exist.
    Supports EXCEL_STORAGE_DIR environment variable override and DEFAULT_EXCEL_DIR patching.
    """
    override = os.getenv("EXCEL_STORAGE_DIR", "").strip()
    if override:
        excel_dir = Path(override).expanduser().resolve()
    else:
        excel_dir = Path(DEFAULT_EXCEL_DIR).resolve()

    excel_dir.mkdir(parents=True, exist_ok=True)
    return excel_dir


def resolve_excel_path(
    filename: str,
    allow_nonexistent: bool = False,
    excel_dir: Path | None = None,
) -> Path:
    """
    Safely resolves a filename within data/excel/, strictly jailing against
    directory traversal attacks (.., absolute paths outside data/excel/, null bytes).
    Enforces .xlsx file extension (auto-appends if omitted).

    Args:
        filename: Name or relative subpath of the Excel file.
        allow_nonexistent: If True, does not raise FileNotFoundError if file is absent.
        excel_dir: Optional base directory override (defaults to get_excel_dir()).

    Returns:
        Resolved absolute Path inside the excel storage directory.

    Raises:
        ValueError: If filename is empty, contains null bytes, or has an invalid extension.
        PermissionError: If path escapes the jail directory.
        FileNotFoundError: If allow_nonexistent is False and file does not exist.
    """
    if not isinstance(filename, str):
        raise ValueError("Filename must be a string.")

    clean = filename.strip()
    if not clean:
        raise ValueError("Filename cannot be empty.")
    if "\0" in clean:
        raise ValueError("Filename contains invalid null byte.")

    # Reject bare ".." and "../" or any path component containing ".." before appending .xlsx
    norm_check = clean.replace("\\", "/").strip("/")
    parts = [part for part in norm_check.split("/") if part]
    if norm_check == ".." or any(".." in part for part in parts):
        raise PermissionError(
            f"Access Denied: Path '{filename}' contains invalid parent directory traversal tokens."
        )

    base_dir = (excel_dir or get_excel_dir()).resolve()

    p = Path(clean)
    if p.is_absolute():
        target = p.resolve()
        # Enforce jailing constraint
        try:
            target.relative_to(base_dir)
        except ValueError:
            raise PermissionError(
                f"Access Denied: Path '{filename}' escapes the Excel storage directory."
            )
        if target == base_dir:
            raise ValueError("Target path cannot be the Excel root directory itself.")
        if target.suffix.lower() != ".xlsx":
            raise ValueError(
                f"Invalid file extension '{target.suffix}'. Only '.xlsx' files are supported."
            )
    else:
        # Normalize relative path
        norm = clean.replace("\\", "/").strip("/")
        if norm.startswith("data/excel/"):
            norm = norm[len("data/excel/"):]
        elif norm.startswith("excel/"):
            norm = norm[len("excel/"):]

        candidate = Path(norm)
        if not candidate.suffix:
            candidate = candidate.with_suffix(".xlsx")
        elif candidate.suffix.lower() != ".xlsx":
            raise ValueError(
                f"Invalid file extension '{candidate.suffix}'. Only '.xlsx' files are supported."
            )

        target = (base_dir / candidate).resolve()
        try:
            target.relative_to(base_dir)
        except ValueError:
            raise PermissionError(
                f"Access Denied: Path '{filename}' escapes the Excel storage directory."
            )
        if target == base_dir:
            raise ValueError("Target path cannot be the Excel root directory itself.")

    if not allow_nonexistent and not target.exists():
        raise FileNotFoundError(f"Excel file not found: '{filename}'")

    return target


# =========================================================================
# Polymorphic Workbook Context & Helpers
# =========================================================================

@contextmanager
def open_workbook(
    wb_or_file: WorkbookInput,
    save_on_exit: bool = False,
    data_only: bool = False,
) -> Generator[tuple[Workbook, bool]]:
    """
    Context manager yielding (workbook, is_file).
    If wb_or_file is a file path/name and save_on_exit is True, saves changes upon exit.
    Guarantees workbook closure when loaded from disk.
    """
    if isinstance(wb_or_file, Workbook):
        yield wb_or_file, False
    elif isinstance(wb_or_file, (str, Path)):
        wb = load_workbook(wb_or_file, data_only=data_only)
        try:
            yield wb, True
            if save_on_exit:
                save_workbook(wb, wb_or_file)
        finally:
            wb.close()
    else:
        raise TypeError(f"Expected Workbook, str, or Path, got {type(wb_or_file).__name__}")


def get_sheet(wb: Workbook, sheet_name: str | None = None) -> Worksheet:
    """
    Retrieves worksheet by exact or case-insensitive title, defaulting to active sheet.
    Raises ValueError if sheet is not found.
    """
    if sheet_name is None:
        if wb.active is not None:
            return wb.active
        elif wb.worksheets:
            return wb.worksheets[0]
        else:
            raise ValueError("Workbook contains no worksheets.")

    clean_name = str(sheet_name).strip()
    if clean_name in wb.sheetnames:
        return wb[clean_name]

    # Case-insensitive fallback
    sheet_map = {name.lower(): name for name in wb.sheetnames}
    if clean_name.lower() in sheet_map:
        return wb[sheet_map[clean_name.lower()]]

    raise ValueError(f"Sheet '{sheet_name}' not found. Available sheets: {wb.sheetnames}")


# =========================================================================
# Workbook CRUD & In-Memory Serialization
# =========================================================================

TEMPLATES: dict[str, dict[str, Any]] = {
    "blank": {
        "sheet_name": "Sheet1",
        "headers": [],
        "rows": [],
    },
    "budget": {
        "sheet_name": "Budget",
        "headers": ["Item", "Category", "Planned", "Actual", "Difference"],
        "rows": [
            ["Rent", "Housing", 1200, 1200, "=C2-D2"],
            ["Groceries", "Food", 400, 450, "=C3-D3"],
            ["Utilities", "Utilities", 150, 130, "=C4-D4"],
            ["Internet", "Utilities", 70, 70, "=C5-D5"],
            ["Total", "", "=SUM(C2:C5)", "=SUM(D2:D5)", "=C6-D6"],
        ],
    },
    "project_tracker": {
        "sheet_name": "Tasks",
        "headers": ["Task ID", "Task Description", "Assignee", "Status", "Due Date"],
        "rows": [
            ["TASK-1", "Design Database Schema", "Alex", "Completed", "2026-10-15"],
            ["TASK-2", "Implement API Endpoints", "Sam", "In Progress", "2026-10-20"],
            ["TASK-3", "Write Integration Tests", "Taylor", "Pending", "2026-10-25"],
        ],
    },
    "reading_list": {
        "sheet_name": "Books",
        "headers": ["Title", "Author", "Status", "Rating", "Notes"],
        "rows": [
            ["Clean Architecture", "Robert C. Martin", "Read", 5, "Essential reading"],
            ["Designing Data-Intensive Applications", "Martin Kleppmann", "Reading", 5, "Distributed systems"],
        ],
    },
}


def load_workbook(
    filename_or_path: str | Path | io.BytesIO,
    data_only: bool = False,
) -> openpyxl.Workbook:
    """
    Loads an Excel workbook.
    CRITICAL: data_only MUST always remain False when modifying or preserving
    workbooks to prevent overwriting/destroying formula expressions.
    """
    try:
        if isinstance(filename_or_path, io.BytesIO):
            filename_or_path.seek(0)
            return openpyxl.load_workbook(filename_or_path, data_only=data_only)

        path = resolve_excel_path(str(filename_or_path), allow_nonexistent=False)
        return openpyxl.load_workbook(path, data_only=data_only)
    except (FileNotFoundError, PermissionError):
        raise
    except CORRUPT_EXCEL_EXCEPTIONS as e:
        raise ValueError(f"Invalid or corrupted Excel file: {e}") from e


def save_workbook(
    wb: openpyxl.Workbook,
    filename_or_path: str | Path,
) -> Path:
    """
    Saves an openpyxl Workbook to the sandboxed path.
    Creates parent directories if necessary and returns the resolved Path.
    """
    target = resolve_excel_path(str(filename_or_path), allow_nonexistent=True)
    target.parent.mkdir(parents=True, exist_ok=True)
    wb.save(target)
    return target


def workbook_to_bytes(wb: openpyxl.Workbook) -> bytes:
    """
    Serializes a workbook into an in-memory byte buffer without saving to disk.
    Used for streaming Discord file attachments.
    """
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def get_workbook_bytes(filename: str) -> tuple[str, bytes]:
    """
    Reads a workbook file from data/excel/ and returns its filename and binary payload.
    """
    target = resolve_excel_path(filename, allow_nonexistent=False)
    file_bytes = target.read_bytes()
    return target.name, file_bytes


def create_workbook(
    filename: str,
    sheet_name: str = "Sheet1",
    headers: list[str] | None = None,
    rows: list[list[Any]] | None = None,
    template: str | None = None,
) -> dict[str, Any]:
    """
    Creates a new .xlsx workbook with initial sheet, optional headers, initial rows, or a template.
    Returns metadata dictionary.
    """
    target = resolve_excel_path(filename, allow_nonexistent=True)
    wb = openpyxl.Workbook()
    ws = wb.active

    actual_sheet_name = sheet_name or "Sheet1"
    actual_headers = headers
    actual_rows = rows

    if template:
        norm_template = template.strip().lower()
        if norm_template in TEMPLATES:
            tpl_data = TEMPLATES[norm_template]
            if not sheet_name or sheet_name == "Sheet1":
                actual_sheet_name = tpl_data["sheet_name"]
            if actual_headers is None:
                actual_headers = tpl_data["headers"]
            if actual_rows is None:
                actual_rows = tpl_data["rows"]

    validate_sheet_title(actual_sheet_name)
    ws.title = actual_sheet_name

    current_row = 1
    if actual_headers:
        for col_idx, header in enumerate(actual_headers, start=1):
            coord = f"{get_column_letter(col_idx)}{current_row}"
            write_cell(ws, coord, header)
        current_row += 1

    if actual_rows:
        for row in actual_rows:
            for col_idx, val in enumerate(row, start=1):
                coord = f"{get_column_letter(col_idx)}{current_row}"
                write_cell(ws, coord, val)
            current_row += 1

    save_workbook(wb, target)
    wb.close()

    total_rows = (current_row - 1)
    total_cols = len(actual_headers) if actual_headers else (len(actual_rows[0]) if actual_rows else 0)

    return {
        "filename": target.name,
        "path": str(target),
        "sheet_name": actual_sheet_name,
        "total_rows": total_rows,
        "total_columns": total_cols,
        "status": "created",
    }


# =========================================================================
# Cell & Range Writing Primitives
# =========================================================================

def write_cell(
    target: Worksheet | openpyxl.Workbook | str | Path,
    coordinate: str,
    value: Any,
    sheet_name: str | None = None,
) -> Cell:
    """
    Writes a typed value or formula expression into a worksheet cell.
    Supports Worksheet, Workbook, or filename/Path as target.
    Guarantees that formulas starting with '=' have data_type='f' and
    are preserved without corruption.
    """
    clean_coord = coordinate.strip().upper()
    if not COORD_RE.match(clean_coord):
        raise ValueError(f"Invalid cell coordinate: '{coordinate}'. Expected format like 'A1', 'C10'.")

    if isinstance(target, (str, Path)):
        wb = load_workbook(target, data_only=False)
        try:
            ws = get_sheet(wb, sheet_name)
            cell = _write_cell_internal(ws, clean_coord, value)
            save_workbook(wb, target)
            return cell
        finally:
            wb.close()
    elif isinstance(target, openpyxl.Workbook):
        ws = get_sheet(target, sheet_name)
        return _write_cell_internal(ws, clean_coord, value)
    elif isinstance(target, Worksheet):
        return _write_cell_internal(target, clean_coord, value)
    else:
        raise TypeError(f"Expected Worksheet, Workbook, str, or Path, got {type(target).__name__}")


def _write_cell_internal(ws: Worksheet, clean_coord: str, value: Any) -> Cell:
    """Internal helper to set cell values with formula preservation."""
    cell = ws[clean_coord]

    if value is None:
        cell.value = None
        return cell

    if isinstance(value, str):
        stripped = value.strip()
        if stripped.startswith("="):
            cell.value = stripped
            cell.data_type = "f"
            return cell
        cell.value = value
        return cell

    cell.value = value
    return cell


def write_range(
    target: Worksheet | openpyxl.Workbook | str | Path,
    start_cell: str,
    rows: list[list[Any]],
    sheet_name: str | None = None,
) -> int:
    """
    Writes a 2D matrix of rows starting at start_cell (e.g. 'B2').
    Supports Worksheet, Workbook, or filename/Path as target.
    Returns total number of cells populated.
    """
    clean_start = start_cell.strip().upper()
    if not COORD_RE.match(clean_start):
        raise ValueError(f"Invalid start coordinate: '{start_cell}'.")

    start_row, start_col = coordinate_to_tuple(clean_start)

    if isinstance(target, (str, Path)):
        wb = load_workbook(target, data_only=False)
        try:
            ws = get_sheet(wb, sheet_name)
            cells_written = _write_range_internal(ws, start_row, start_col, rows)
            save_workbook(wb, target)
            return cells_written
        finally:
            wb.close()
    elif isinstance(target, openpyxl.Workbook):
        ws = get_sheet(target, sheet_name)
        return _write_range_internal(ws, start_row, start_col, rows)
    elif isinstance(target, Worksheet):
        return _write_range_internal(target, start_row, start_col, rows)
    else:
        raise TypeError(f"Expected Worksheet, Workbook, str, or Path, got {type(target).__name__}")


def _write_range_internal(
    ws: Worksheet,
    start_row: int,
    start_col: int,
    rows: list[list[Any]],
) -> int:
    cells_written = 0
    for r_offset, row in enumerate(rows):
        for c_offset, val in enumerate(row):
            coord = f"{get_column_letter(start_col + c_offset)}{start_row + r_offset}"
            _write_cell_internal(ws, coord, val)
            cells_written += 1
    return cells_written


def clear_cells(
    wb_or_file: Worksheet | openpyxl.Workbook | str | Path,
    target: str | Sequence[str],
    sheet_name: str | None = None,
) -> int:
    """
    Clears cell values across coordinates, ranges, column/row intervals, or lists of targets.
    Setting cell.value = None resets both content and formula data_type in openpyxl.
    Automatically saves if wb_or_file is a file path.

    Returns:
        Total count of cells cleared.
    """
    if isinstance(wb_or_file, (str, Path)):
        wb = load_workbook(wb_or_file, data_only=False)
        try:
            ws = get_sheet(wb, sheet_name)
            count = _clear_cells_internal(ws, target)
            save_workbook(wb, wb_or_file)
            return count
        finally:
            wb.close()
    elif isinstance(wb_or_file, openpyxl.Workbook):
        ws = get_sheet(wb_or_file, sheet_name)
        return _clear_cells_internal(ws, target)
    elif isinstance(wb_or_file, Worksheet):
        return _clear_cells_internal(wb_or_file, target)
    else:
        raise TypeError(f"Expected Worksheet, Workbook, str, or Path, got {type(wb_or_file).__name__}")


def clear_range(
    target: Worksheet | openpyxl.Workbook | str | Path,
    cell_range: str,
    sheet_name: str | None = None,
) -> int:
    """Alias for clear_cells for range specifications."""
    return clear_cells(target, cell_range, sheet_name=sheet_name)


def _clear_cells_internal(ws: Worksheet, target: str | Sequence[str]) -> int:
    if not target:
        return 0

    if isinstance(target, (list, tuple, set)):
        return sum(_clear_cells_internal(ws, sub_t) for sub_t in target)

    t_str = str(target).strip()
    count = 0
    if ":" in t_str:
        try:
            min_col, min_row, max_col, max_row = range_boundaries(t_str)
        except ValueError:
            raise ValueError(f"Invalid range specification for clearing: '{t_str}'")

        min_col = min_col or 1
        max_col = max_col or ws.max_column
        min_row = min_row or 1
        max_row = max_row or ws.max_row

        for r in range(min_row, max_row + 1):
            for c in range(min_col, max_col + 1):
                cell = ws.cell(row=r, column=c)
                if cell.value is not None:
                    cell.value = None
                    count += 1
    else:
        coord_clean = t_str.upper()
        if not COORD_RE.match(coord_clean):
            raise ValueError(f"Invalid cell coordinate for clearing: '{t_str}'")
        cell = ws[coord_clean]
        if cell.value is not None:
            cell.value = None
            count += 1

    return count


def write_cells(
    filename: str,
    sheet_name: str | None = None,
    updates: list[dict[str, Any]] | None = None,
    start_cell: str | None = None,
    rows: list[list[Any]] | None = None,
    clear_range: str | None = None,
    clear_range_str: str | None = None,
) -> dict[str, Any]:
    """
    Comprehensive cell update handler. Supports:
    1. Clearing a range (clear_range or clear_range_str)
    2. Writing a 2D matrix starting at start_cell (start_cell + rows)
    3. Writing individual coordinate updates (updates=[{"cell": "A1", "value": 10}, ...])
    """
    wb = load_workbook(filename, data_only=False)
    try:
        ws = get_sheet(wb, sheet_name)

        target_clear = clear_range or clear_range_str
        cleared_count = 0
        if target_clear:
            cleared_count = clear_cells(ws, target_clear)

        matrix_cells = 0
        if start_cell and rows:
            matrix_cells = write_range(ws, start_cell, rows)

        individual_updates = 0
        if updates:
            for item in updates:
                coord = item.get("cell") or item.get("coordinate")
                val = item.get("value")
                if not coord:
                    continue
                write_cell(ws, str(coord), val)
                individual_updates += 1

        save_workbook(wb, filename)

        return {
            "filename": filename,
            "sheet_name": ws.title,
            "cleared_cells": cleared_count,
            "matrix_cells_written": matrix_cells,
            "individual_cells_updated": individual_updates,
            "total_modifications": cleared_count + matrix_cells + individual_updates,
            "status": "updated",
        }
    finally:
        wb.close()


# =========================================================================
# Sheet Lifecycle Management
# =========================================================================

def validate_sheet_title(
    title: str,
    existing_sheets: Sequence[str] | None = None,
    current_title: str | None = None,
) -> str:
    """
    Strictly validates sheet title according to Excel and openpyxl rules:
    - Must be a non-empty string.
    - Maximum 31 characters.
    - Cannot contain forbidden characters: : \\ / ? * [ ]
    - Cannot start or end with an apostrophe (').
    - Cannot be named 'History' (reserved by Excel).
    - Cannot duplicate an existing sheet name (case-insensitive).
    """
    if not isinstance(title, str):
        raise TypeError(f"Sheet title must be a string, got {type(title).__name__}")

    clean = title.strip()
    if not clean:
        raise ValueError("Sheet title cannot be empty or whitespace only.")

    if len(clean) > MAX_SHEET_TITLE_LEN:
        raise ValueError(
            f"Sheet title '{clean}' exceeds Excel's 31-character limit ({len(clean)} characters)."
        )

    found_forbidden = [c for c in clean if c in FORBIDDEN_SHEET_CHARS]
    if found_forbidden:
        unique_bad = sorted(set(found_forbidden))
        raise ValueError(
            f"Sheet title '{clean}' contains forbidden characters: {unique_bad}. "
            "Characters : \\ / ? * [ ] are not allowed in Excel sheet names."
        )

    if clean.startswith("'") or clean.endswith("'"):
        raise ValueError("Sheet title cannot start or end with an apostrophe/single quote (').")

    if clean.lower() == "history":
        raise ValueError("Sheet title cannot be 'History' (reserved by Excel for revision tracking).")

    if existing_sheets:
        target_lower = clean.lower()
        for existing in existing_sheets:
            if current_title and existing.lower() == current_title.lower():
                continue
            if existing.lower() == target_lower:
                raise ValueError(
                    f"Sheet '{clean}' already exists in workbook (case-insensitive conflict with '{existing}')."
                )

    return clean


def list_sheets(wb_or_file: WorkbookInput) -> list[str]:
    """
    Returns an ordered list of all worksheet titles in the workbook.
    """
    with open_workbook(wb_or_file, save_on_exit=False) as (wb, _):
        return list(wb.sheetnames)


def add_sheet(
    wb_or_file: WorkbookInput,
    title: str,
    index: int | None = None,
) -> Worksheet:
    """
    Adds a new worksheet to the workbook, enforcing Excel constraints.
    Automatically saves if wb_or_file is a file path.
    """
    with open_workbook(wb_or_file, save_on_exit=True) as (wb, _):
        clean_title = validate_sheet_title(title, existing_sheets=wb.sheetnames)
        ws = wb.create_sheet(title=clean_title, index=index)
        return ws


def delete_sheet(wb_or_file: WorkbookInput, title: str) -> None:
    """
    Deletes a worksheet by name, enforcing the single-sheet workbook guard:
    - Cannot delete the only worksheet in a workbook.
    - Resets active sheet if deleted sheet was active.
    - Automatically saves if wb_or_file is a file path.
    """
    with open_workbook(wb_or_file, save_on_exit=True) as (wb, _):
        if len(wb.sheetnames) <= 1:
            raise ValueError(
                "Cannot delete the only sheet in a workbook. Excel workbooks must contain at least one sheet."
            )

        ws = get_sheet(wb, title)
        was_active = (wb.active == ws)
        wb.remove(ws)

        if (was_active or wb.active is None) and wb.worksheets:
            wb.active = wb.worksheets[0]


def rename_sheet(
    wb_or_file: WorkbookInput,
    old_title: str,
    new_title: str,
) -> Worksheet:
    """
    Renames an existing worksheet, validating new title constraints and uniqueness.
    Automatically saves if wb_or_file is a file path.
    """
    with open_workbook(wb_or_file, save_on_exit=True) as (wb, _):
        ws = get_sheet(wb, old_title)
        clean_new = validate_sheet_title(
            new_title,
            existing_sheets=wb.sheetnames,
            current_title=ws.title,
        )
        if ws.title.lower() == clean_new.lower() and ws.title != clean_new:
            ws.title = f"__tmp_ren_{int(time.time() * 1000)}__"
        ws.title = clean_new
        return ws


def set_active_sheet(wb_or_file: WorkbookInput, title: str) -> Worksheet:
    """
    Sets the active worksheet by title.
    Automatically saves if wb_or_file is a file path.
    """
    with open_workbook(wb_or_file, save_on_exit=True) as (wb, _):
        ws = get_sheet(wb, title)
        wb.active = ws
        return ws


def manage_sheet(
    filename: str,
    action: str,
    sheet_name: str | None = None,
    new_name: str | None = None,
) -> dict[str, Any]:
    """
    High-level sheet management dispatcher for AI tools layer:
    Actions: 'list', 'add', 'rename', 'delete', 'set_active'.
    Returns structured status dictionary.
    """
    clean_action = str(action or "").strip().lower()

    if clean_action == "list":
        sheets = list_sheets(filename)
        with open_workbook(filename) as (wb, _):
            active_sheet = wb.active.title if wb.active else (sheets[0] if sheets else None)
        return {
            "action": "list",
            "filename": filename,
            "sheets": sheets,
            "active_sheet": active_sheet,
            "total_sheets": len(sheets),
        }

    elif clean_action == "add":
        if not sheet_name:
            raise ValueError("Parameter 'sheet_name' is required for action 'add'.")
        add_sheet(filename, title=sheet_name)
        sheets = list_sheets(filename)
        return {
            "action": "add",
            "filename": filename,
            "sheet_name": sheet_name,
            "sheets": sheets,
            "message": f"Added worksheet '{sheet_name}'.",
        }

    elif clean_action == "rename":
        if not sheet_name or not new_name:
            raise ValueError("Parameters 'sheet_name' (old) and 'new_name' are required for action 'rename'.")
        rename_sheet(filename, old_title=sheet_name, new_title=new_name)
        sheets = list_sheets(filename)
        return {
            "action": "rename",
            "filename": filename,
            "old_name": sheet_name,
            "new_name": new_name,
            "sheets": sheets,
            "message": f"Renamed worksheet '{sheet_name}' to '{new_name}'.",
        }

    elif clean_action == "delete":
        if not sheet_name:
            raise ValueError("Parameter 'sheet_name' is required for action 'delete'.")
        delete_sheet(filename, title=sheet_name)
        sheets = list_sheets(filename)
        with open_workbook(filename) as (wb, _):
            active_sheet = wb.active.title if wb.active else (sheets[0] if sheets else None)
        return {
            "action": "delete",
            "filename": filename,
            "deleted_sheet": sheet_name,
            "sheets": sheets,
            "active_sheet": active_sheet,
            "message": f"Deleted worksheet '{sheet_name}'.",
        }

    elif clean_action == "set_active":
        if not sheet_name:
            raise ValueError("Parameter 'sheet_name' is required for action 'set_active'.")
        ws = set_active_sheet(filename, title=sheet_name)
        return {
            "action": "set_active",
            "filename": filename,
            "active_sheet": ws.title,
            "message": f"Set active worksheet to '{ws.title}'.",
        }

    else:
        raise ValueError(
            f"Unknown sheet action '{action}'. Valid actions: 'list', 'add', 'rename', 'delete', 'set_active'."
        )


# =========================================================================
# Cell & Range Reading
# =========================================================================

def read_cell(
    wb_or_file: WorkbookInput,
    coordinate: str,
    sheet_name: str | None = None,
) -> dict[str, Any]:
    """
    Reads a single cell by coordinate (e.g. 'A1', 'B12').
    Returns structured dictionary with coordinate, value, formula, data_type, and formatted text.
    """
    coord_clean = str(coordinate or "").strip().upper()
    if not COORD_RE.match(coord_clean):
        raise ValueError(
            f"Invalid cell coordinate '{coordinate}'. Expected coordinate format such as 'A1', 'B10'."
        )

    with open_workbook(wb_or_file, save_on_exit=False, data_only=False) as (wb, _):
        ws = get_sheet(wb, sheet_name)
        cell = ws[coord_clean]
        val = cell.value
        data_type = cell.data_type

        is_formula = data_type == "f" or (isinstance(val, str) and val.startswith("="))
        formula = str(val) if is_formula else None

        if val is None:
            formatted = ""
        elif hasattr(val, "isoformat"):
            formatted = val.isoformat()
        else:
            formatted = str(val)

        return {
            "coordinate": coord_clean,
            "sheet_name": ws.title,
            "value": val,
            "formatted": formatted,
            "data_type": data_type,
            "is_formula": is_formula,
            "formula": formula,
        }


def read_range(
    filename: str | Path | openpyxl.Workbook | None = None,
    sheet_name: str | None = None,
    cell_range: str | None = None,
    max_rows: int = 50,
    show_formulas: bool = False,
    has_headers: bool | None = None,
    filename_or_wb: str | Path | openpyxl.Workbook | None = None,
) -> dict[str, Any]:
    """
    Reads a rectangular range or full worksheet, returning an aligned Markdown table
    and raw grid values.
    """
    target = filename if filename is not None else filename_or_wb
    if target is None:
        raise ValueError("A filename or workbook instance must be provided to read_range.")

    with open_workbook(target, save_on_exit=False, data_only=False) as (wb, _):
        # Heuristic: if sheet_name looks like a cell range (contains ':' or matches coordinate)
        # and is not an actual sheet name, treat it as cell_range
        actual_sheet = sheet_name
        actual_range = cell_range
        if sheet_name and (":" in sheet_name or COORD_RE.match(sheet_name.upper())) and sheet_name not in wb.sheetnames:
            actual_range = sheet_name
            actual_sheet = None

        ws = get_sheet(wb, actual_sheet)

        # Detect empty sheet
        is_empty = (ws.max_row == 0) or (
            ws.max_row == 1 and ws.max_column == 1 and ws.cell(1, 1).value is None
        )
        if is_empty:
            return {
                "sheet_name": ws.title,
                "range": actual_range or "A1:A1",
                "table_markdown": "*(empty sheet)*",
                "total_rows": 0,
                "truncated": False,
                "headers": [],
                "rows": [],
            }

        # Resolve range boundaries
        if actual_range:
            clean_range = actual_range.strip().upper()
            try:
                min_col, min_row, max_col, max_row = range_boundaries(clean_range)
            except ValueError:
                raise ValueError(f"Invalid range specification: '{actual_range}'.")
            min_col = min_col or 1
            max_col = max_col or ws.max_column
            min_row = min_row or 1
            max_row = max_row or ws.max_row
        else:
            min_col, min_row = 1, 1
            max_col, max_row = ws.max_column, ws.max_row
            clean_range = f"{get_column_letter(min_col)}{min_row}:{get_column_letter(max_col)}{max_row}"

        raw_rows: list[list[Any]] = []
        for r in range(min_row, max_row + 1):
            row_vals: list[Any] = []
            for c in range(min_col, max_col + 1):
                cell = ws.cell(row=r, column=c)
                v = cell.value
                if show_formulas and (cell.data_type == "f" or (isinstance(v, str) and v.startswith("="))):
                    v = str(v)
                row_vals.append(v)
            raw_rows.append(row_vals)

        if not raw_rows:
            return {
                "sheet_name": ws.title,
                "range": clean_range,
                "table_markdown": "*(no data)*",
                "total_rows": 0,
                "truncated": False,
                "headers": [],
                "rows": [],
            }

        col_letters = [get_column_letter(c) for c in range(min_col, max_col + 1)]

        if has_headers is True:
            headers = [str(c if c is not None else "") for c in raw_rows[0]]
            data_rows = raw_rows[1:]
        elif has_headers is False:
            headers = col_letters
            data_rows = raw_rows
        else:
            # Auto-detect headers: row 1 starts at 1, contains non-empty strings, and has more rows
            first_row = raw_rows[0]
            if (
                min_row == 1
                and len(raw_rows) > 1
                and any(c is not None and str(c).strip() for c in first_row)
                and all(isinstance(c, str) or c is None for c in first_row)
            ):
                headers = [str(c if c is not None else "") for c in first_row]
                data_rows = raw_rows[1:]
            else:
                headers = col_letters
                data_rows = raw_rows

        total_rows = len(data_rows)
        row_truncated = False
        if total_rows > max_rows:
            data_rows = data_rows[:max_rows]
            row_truncated = True

        table_md = format_markdown_table(data_rows, headers=headers, max_chars=1900)
        is_truncated = row_truncated or ("truncated for Discord display" in table_md)

        return {
            "sheet_name": ws.title,
            "range": clean_range,
            "table_markdown": table_md,
            "total_rows": total_rows,
            "truncated": is_truncated,
            "headers": headers,
            "rows": data_rows,
        }


# =========================================================================
# Markdown Table Formatting with Discord Truncation
# =========================================================================

def format_markdown_table(
    rows: Sequence[Sequence[Any]],
    headers: bool | Sequence[str] | None = None,
    max_chars: int = 1900,
    max_table_chars: int | None = None,
    max_cell_len: int = 60,
    code_block: bool = False,
) -> str:
    """
    Renders a 2D grid into an aligned Markdown table with pipe delimiters, space padding,
    left-aligned separators (|:---|:------|), and automatic row-level truncation strictly
    respecting Discord's message cap.

    Args:
        rows: 2D matrix of cell values.
        headers: True to use rows[0] as headers, or sequence of header strings, or None/False.
        max_chars: Upper character limit for output (default 1900).
        max_table_chars: Optional override for max_chars.
        max_cell_len: Truncates individual cells exceeding this length with '...' (default 60).
        code_block: If True, wraps table in monospace code fence.
    """
    limit = max_table_chars if max_table_chars is not None else max_chars

    if not rows and not headers:
        return "*(empty table)*"

    def clean_cell(val: Any) -> str:
        if val is None:
            return ""
        if hasattr(val, "isoformat"):
            s = val.isoformat()
        else:
            s = str(val)
        s = s.replace("\r\n", " ").replace("\n", " ").replace("\r", " ")
        s = s.replace("|", "\\|")
        if len(s) > max_cell_len:
            s = s[: max_cell_len - 3] + "..."
        return s

    if headers is True:
        if rows:
            clean_headers = [clean_cell(h) for h in rows[0]]
            clean_rows = [[clean_cell(c) for c in r] for r in rows[1:]]
        else:
            clean_headers = []
            clean_rows = []
    elif isinstance(headers, (list, tuple)):
        clean_headers = [clean_cell(h) for h in headers]
        clean_rows = [[clean_cell(c) for c in r] for r in rows]
    else:
        clean_headers = []
        clean_rows = [[clean_cell(c) for c in r] for r in rows]

    num_cols = len(clean_headers) if clean_headers else (max(len(r) for r in clean_rows) if clean_rows else 0)
    if num_cols == 0:
        return "*(empty table)*"

    # Pad ragged rows to uniform column count
    for r in clean_rows:
        if len(r) < num_cols:
            r.extend([""] * (num_cols - len(r)))
        elif len(r) > num_cols:
            r = r[:num_cols]

    # Calculate column widths (minimum width 1)
    col_widths = [1] * num_cols
    if clean_headers:
        for i, h in enumerate(clean_headers):
            col_widths[i] = max(col_widths[i], len(h))
    else:
        clean_headers = [get_column_letter(i + 1) for i in range(num_cols)]
        for i, h in enumerate(clean_headers):
            col_widths[i] = max(col_widths[i], len(h))

    for r in clean_rows:
        for i in range(num_cols):
            col_widths[i] = max(col_widths[i], len(r[i]))

    # Build header line and left-aligned separator line
    budget = limit - (8 if code_block else 0)

    # Enforce column truncation if total header line exceeds character budget.
    # In markdown tables, header line and separator line have identical character length.
    # We allocate character budget so that header + sep + optional rows/truncation notice fit within budget.
    header_budget = max(50, (budget - (120 if clean_rows else 20)) // (3 if clean_rows else 2))

    full_header_line = "| " + " | ".join(clean_headers[i].ljust(col_widths[i]) for i in range(num_cols)) + " |"
    if (len(full_header_line) > header_budget or (2 * len(full_header_line) + 2) > budget) and num_cols > 1:
        target_k = 1
        for k in range(num_cols - 1, 0, -1):
            omission_hdr = f"... (+{num_cols - k} columns omitted)"
            omission_w = len(omission_hdr)
            cand_widths = [col_widths[i] for i in range(k)] + [omission_w]
            cand_headers = [clean_headers[i] for i in range(k)] + [omission_hdr]
            cand_h = "| " + " | ".join(cand_headers[i].ljust(cand_widths[i]) for i in range(k + 1)) + " |"
            if len(cand_h) <= header_budget and (2 * len(cand_h) + 2) <= budget:
                target_k = k
                break

        N = num_cols - target_k
        omission_hdr = f"... (+{N} columns omitted)"
        clean_headers = clean_headers[:target_k] + [omission_hdr]
        clean_rows = [r[:target_k] + ["..."] for r in clean_rows]
        num_cols = target_k + 1
        col_widths = col_widths[:target_k] + [len(omission_hdr)]

    lines: list[str] = []
    header_line = "| " + " | ".join(clean_headers[i].ljust(col_widths[i]) for i in range(num_cols)) + " |"
    sep_line = "|" + "|".join(":" + "-" * (col_widths[i] + 1) for i in range(num_cols)) + "|"
    lines.append(header_line)
    lines.append(sep_line)

    if not clean_rows:
        res = "\n".join(lines)
        if len(res) > limit:
            res = res[:limit]
        return f"```\n{res}\n```" if code_block else res

    current_char_count = sum(len(line) + 1 for line in lines)
    displayed_rows = 0

    for idx, r in enumerate(clean_rows):
        row_line = "| " + " | ".join(r[i].ljust(col_widths[i]) for i in range(num_cols)) + " |"
        trunc_msg = f"*... [truncated for Discord display: showing {displayed_rows} of {len(clean_rows)} rows]*"

        if current_char_count + len(row_line) + 1 + len(trunc_msg) + 1 > budget:
            lines.append(f"*... [truncated for Discord display: showing {displayed_rows} of {len(clean_rows)} rows]*")
            break

        lines.append(row_line)
        current_char_count += len(row_line) + 1
        displayed_rows += 1

    res = "\n".join(lines)
    if code_block:
        res = f"```\n{res}\n```"

    # Invariant: Output must NEVER exceed limit
    if len(res) > limit:
        cutoff = limit - len("\n*... [truncated]*")
        if code_block:
            cutoff = limit - len("```\n\n*... [truncated]*\n```")
        if cutoff > 0:
            res = res[:cutoff] + "\n*... [truncated]*"
            if code_block:
                res = f"```\n{res}\n```"
        else:
            res = res[:limit]

    return res


# =========================================================================
# AI Summary Extraction & Storage Purging
# =========================================================================

def get_workbook_summary(
    filename_or_bytes: str | Path | bytes,
    max_rows_preview: int = 5,
    max_cols_preview: int = 10,
    max_sheets_preview: int = 5,
) -> str:
    """
    Extract structural metadata and a header preview from an Excel workbook
    formatted as Markdown text for Gemini chat context injection.

    Args:
        filename_or_bytes: Either a filename/Path in data/excel/ or raw bytes.
        max_rows_preview: Reserved for sample rows (default 5).
        max_cols_preview: Max columns to show in header preview (default 10).
        max_sheets_preview: Max sheets to detail in summary (default 5).

    Returns:
        Structured Markdown summary of worksheets, dimensions, and headers.

    Raises:
        FileNotFoundError: If the specified file does not exist.
        ValueError: If file is not a valid OpenXML Excel workbook.
    """
    display_name = "Uploaded Workbook"
    wb: openpyxl.Workbook | None = None

    try:
        if isinstance(filename_or_bytes, bytes):
            if not zipfile.is_zipfile(io.BytesIO(filename_or_bytes)):
                raise ValueError("The provided file is not a valid OpenXML (.xlsx) archive.")
            wb = openpyxl.load_workbook(io.BytesIO(filename_or_bytes), data_only=False)
        else:
            resolved_path = resolve_excel_path(str(filename_or_bytes), allow_nonexistent=False)
            display_name = resolved_path.name
            wb = openpyxl.load_workbook(resolved_path, data_only=False)

        sheet_names = wb.sheetnames
        if not sheet_names:
            return f"### Excel Workbook: `{display_name}`\n*(Empty workbook with no sheets)*"

        active_title = wb.active.title if wb.active else sheet_names[0]

        lines = [
            f"### Excel Workbook: `{display_name}`",
            f"- **Worksheets ({len(sheet_names)})**: " + ", ".join(
                f"`{name}`" + (" (Active)" if name == active_title else "") for name in sheet_names
            ),
            f"- **Active Sheet**: `{active_title}`",
            "",
        ]

        detailed_sheets = sheet_names[:max_sheets_preview]
        for name in detailed_sheets:
            ws = wb[name]
            is_active = (name == active_title)
            active_marker = " (Active)" if is_active else ""
            lines.append(f"#### Sheet: `{name}`{active_marker}")

            max_row = ws.max_row or 0
            max_col = ws.max_column or 0

            is_empty = False
            if max_row <= 1 and max_col <= 1:
                val = ws.cell(row=1, column=1).value
                if val is None or str(val).strip() == "":
                    is_empty = True

            if is_empty or max_row == 0 or max_col == 0:
                lines.append("- **Status**: Empty worksheet (0 rows × 0 columns)")
                lines.append("")
                continue

            dimensions = ws.dimensions or f"A1:{get_column_letter(max_col)}{max_row}"
            lines.append(f"- **Dimensions**: `{dimensions}` ({max_row} rows × {max_col} columns)")

            col_limit = min(max_col, max_cols_preview)
            headers: list[str] = []
            for col_idx in range(1, col_limit + 1):
                cell_val = ws.cell(row=1, column=col_idx).value
                if cell_val is None or str(cell_val).strip() == "":
                    headers.append("(empty)")
                else:
                    cleaned = str(cell_val).replace("\n", " ").replace("|", "\\|").strip()
                    headers.append(f"`{cleaned}`")

            header_str = " | ".join(headers)
            if max_col > max_cols_preview:
                header_str += f" | *(+{max_col - max_cols_preview} more columns)*"

            lines.append(f"- **Headers (Row 1)**: {header_str}")
            lines.append("")

        if len(sheet_names) > max_sheets_preview:
            lines.append(f"*(+{len(sheet_names) - max_sheets_preview} additional sheets not previewed)*\n")

        lines.append(
            "*Tip: Use spreadsheet tools (`read_spreadsheet_range`, `update_spreadsheet_cells`, `manage_spreadsheet_sheets`) to query or modify this workbook.*"
        )
        return "\n".join(lines).strip()

    except (FileNotFoundError, PermissionError):
        raise
    except CORRUPT_EXCEL_EXCEPTIONS as e:
        raise ValueError(f"Invalid or corrupted Excel file: {e}") from e
    finally:
        if wb is not None:
            wb.close()


def purge_old_files(
    max_age_seconds: int = 7 * 86400,
    directory: Path | str | None = None,
) -> list[str]:
    """
    Scans the Excel storage directory and removes files older than max_age_seconds.

    Protects against directory traversal and symlink attacks by verifying all
    candidate paths remain strictly within the target storage directory.

    Args:
        max_age_seconds: Maximum allowed file age in seconds based on mtime.
                         Defaults to 7 days. If <= 0, purges all regular files.
        directory: Optional directory path to scan. Defaults to get_excel_dir().

    Returns:
        List of deleted file names (basenames).
    """
    if directory is None:
        target_dir = get_excel_dir()
    else:
        target_dir = Path(directory).resolve()

    if not target_dir.exists() or not target_dir.is_dir():
        logger.debug(f"Excel directory '{target_dir}' does not exist; skipping purge.")
        return []

    now = time.time()
    deleted_files: list[str] = []

    try:
        entries = list(target_dir.iterdir())
    except OSError as e:
        logger.error(f"Failed to list Excel directory '{target_dir}': {e}")
        return []

    for entry in entries:
        try:
            if entry.is_symlink():
                resolved = entry.resolve()
                if not resolved.is_relative_to(target_dir):
                    entry.unlink(missing_ok=True)
                    deleted_files.append(entry.name)
                    continue

            if not entry.is_file():
                continue

            if not entry.resolve().is_relative_to(target_dir):
                continue

            stat = entry.stat(follow_symlinks=False)
            file_age = now - stat.st_mtime

            if max_age_seconds <= 0 or file_age >= max_age_seconds:
                entry.unlink()
                deleted_files.append(entry.name)
                logger.info(f"Purged expired Excel file: {entry.name} (age: {int(file_age)}s)")

        except FileNotFoundError:
            continue
        except OSError as e:
            logger.warning(f"Could not purge file '{entry.name}': {e}")

    return deleted_files
